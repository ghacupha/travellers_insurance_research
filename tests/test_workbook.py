# SPDX-FileCopyrightText: 2026 ghacupha
# SPDX-License-Identifier: MIT

"""Executable workbook contract, including saved formula caches and live links."""
from pathlib import Path
import tempfile
import unittest

try:
    from openpyxl import load_workbook
    from bizplan.insurance.workbook import build_workbook
except ImportError:
    load_workbook = build_workbook = None


ROOT = Path(__file__).resolve().parents[1]


@unittest.skipIf(build_workbook is None, 'install release dependencies and openpyxl')
class WorkbookTests(unittest.TestCase):
    def test_scenario_assumption_model_and_output_are_formula_linked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = build_workbook(ROOT, Path(directory) / 'model.xlsx')
            formulas = load_workbook(path, read_only=False, data_only=False)
            cached = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(formulas.sheetnames,
                             ['Cover', 'Summary', 'Assumptions', 'Scenarios',
                              'Model', 'Output', 'Sources'])
            selector = formulas['Scenarios']['D6']
            self.assertEqual(selector.value, 'Base')
            self.assertEqual(formulas['Scenarios'].data_validations.dataValidation[0].sqref,
                             'D6')
            self.assertIn('Scenarios!$D$6', formulas['Assumptions']['N9'].value)
            self.assertIn('Assumptions!N9', formulas['Model']['N13'].value)
            self.assertEqual(formulas['Output']['D11'].value, '=Model!R22')
            self.assertEqual(formulas['Output']['D13'].value, '=Summary!R15')
            self.assertGreater(cached['Output']['D11'].value, 0)
            self.assertGreater(cached['Output']['D12'].value, 0)
            self.assertEqual(cached['Output']['D10'].value, 'Base')
            self.assertEqual(cached['Output']['D14'].value,
                             'OPEN: historical stock bridges')


if __name__ == '__main__':
    unittest.main()
