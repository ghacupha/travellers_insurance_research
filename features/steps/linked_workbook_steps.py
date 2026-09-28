# SPDX-FileCopyrightText: 2026 ghacupha
# SPDX-License-Identifier: MIT

from pathlib import Path
import tempfile

from behave import given, then
from openpyxl import load_workbook

from bizplan.insurance.workbook import build_workbook


ROOT = Path(__file__).resolve().parents[2]


@given('the Python Travelers workbook is built')
def build(context):
    temporary = tempfile.TemporaryDirectory()
    context.add_cleanup(temporary.cleanup)
    path = build_workbook(ROOT, Path(temporary.name) / 'travelers.xlsx')
    context.workbook_formulas = load_workbook(path, data_only=False)
    context.workbook_values = load_workbook(path, data_only=True)


@then('its scenario selector offers Base, Upside and Downside')
def selector(context):
    sheet = context.workbook_formulas['Scenarios']
    assert sheet['D6'].value == 'Base'
    assert sheet.data_validations.dataValidation[0].formula1 == '"Base,Upside,Downside"'


@then('its active assumptions feed the forecast schedules')
def assumptions(context):
    book = context.workbook_formulas
    assert 'Scenarios!$D$6' in book['Assumptions']['N9'].value
    assert 'Assumptions!N9' in book['Model']['N13'].value


@then('its final research output links to the selected forecast')
def output(context):
    sheet = context.workbook_formulas['Output']
    assert sheet['D10'].value == '=Scenarios!D6'
    assert sheet['D11'].value == '=Model!R22'
    assert sheet['D12'].value.startswith('=Model!R')
    assert sheet['D13'].value == '=Summary!R15'


@then('its saved Base case contains calculated output values')
def cache(context):
    sheet = context.workbook_values['Output']
    assert sheet['D10'].value == 'Base'
    assert sheet['D11'].value > 0
    assert sheet['D12'].value > 0
    assert 0 < sheet['D13'].value < 2
